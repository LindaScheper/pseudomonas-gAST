"""Neue Tabelle blockweise einlesen (ein Aufruf verarbeitet mehrere Blaetter)."""
import openpyxl, pandas as pd, numpy as np, pickle, os, sys, json, time
F="Mutationen_pro_Gen_mit_value_type.xlsx"   # liegt im Arbeitsverzeichnis
AB=["MEM","CAZ","CIP","TOB","GEN"]
STATE="v3_state.json"
st=json.load(open(STATE)) if os.path.exists(STATE) else {"done":[]}
wb=openpyxl.load_workbook(F,read_only=True,data_only=True)
todo=[g for g in wb.sheetnames if not g.startswith("_") and g not in st["done"]]
if not todo:
    print("alle Blaetter gelesen"); wb.close(); sys.exit()
rows=[]; pheno={}
t0=time.time()
for g in todo:
    if time.time()-t0>85: break
    ws=wb[g]; it=ws.iter_rows(values_only=True); hdr=[str(c) for c in next(it)]
    ix={h:i for i,h in enumerate(hdr)}
    for r in it:
        if r is None or r[0] is None: continue
        iso=str(r[0]).strip()
        def nm(k):
            v=r[ix[k]] if k in ix and ix[k]<len(r) else None
            try: return float(str(v).strip())
            except: return np.nan
        rows.append((g,iso,"" if r[1] is None else str(r[1]).strip(),nm("proteinlength"),nm("ref_length")))
        if iso not in pheno:
            d={}
            for ab in AB:
                d[f"v_{ab}"]=nm(f"value_{ab}"); d[f"n_{ab}"]=nm(f"norm_value_{ab}")
                d[f"bp_{ab}"]=nm(f"breakpoint_{ab}")
                t=r[ix[f"type_{ab}"]] if f"type_{ab}" in ix else None
                d[f"t_{ab}"]="" if t is None else str(t).strip()
            pheno[iso]=d
    st["done"].append(g)
wb.close()
M=pd.DataFrame(rows,columns=["gene","isolate","mut_str","plen","reflen"])
M.to_pickle(f"v3_part_{len(st['done'])}.pkl")
pickle.dump(pheno,open(f"v3_ph_{len(st['done'])}.pkl","wb"))
json.dump(st,open(STATE,"w"))
print(f"{len(st['done'])}/96 Blaetter, dieser Lauf {len(rows):,} Zeilen, {len(pheno):,} neue Isolate")
